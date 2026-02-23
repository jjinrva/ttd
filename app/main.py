from __future__ import annotations

import asyncio
import csv
import json
import logging
from logging.handlers import RotatingFileHandler
import os
from pathlib import Path
import textwrap
from datetime import datetime
from io import BytesIO
from typing import Any, Literal
from uuid import uuid4

import httpx
import orjson
from bs4 import BeautifulSoup
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, Field
from pydantic_settings import BaseSettings
from pypdf import PdfReader
from readability import Document as ReadabilityDocument
from PIL import Image, ImageColor, ImageDraw, ImageEnhance, ImageFont


class AppSettings(BaseSettings):
    app_id: str = "meme_gen"
    output_dir: str = "/mnt/outputs"
    data_dir: str = "/data"
    models_dir: str = "/mnt/models"
    max_fetch_mb: int = 100


settings = AppSettings()
DATA_DIR = Path(settings.data_dir)
OUTPUT_DIR = Path(settings.output_dir)
MODELS_DIR = Path(settings.models_dir)
LOG_DIR = DATA_DIR / "logs"
FONTS_DIR = DATA_DIR / "fonts"

for path in [DATA_DIR, OUTPUT_DIR, LOG_DIR, FONTS_DIR]:
    path.mkdir(parents=True, exist_ok=True)

logger = logging.getLogger("meme_gen")
logger.setLevel(logging.INFO)
handler = RotatingFileHandler(LOG_DIR / "app.log", maxBytes=1_000_000, backupCount=3)
handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
logger.addHandler(handler)


DEFAULT_TONES = {
    "funny": "You are a funny comedian. Summarize the content in playful meme language.",
    "concise": "You are concise. Produce a short, punchy meme caption.",
    "serious": "You are serious. Make a factual meme caption.",
    "sarcastic": "You are sarcastic. Generate a witty sarcastic meme caption.",
    "absurd": "You are absurd. Create surreal and absurd meme text.",
    "wholesome": "You are wholesome. Make warm supportive meme text.",
}

DEFAULT_APP_SETTINGS = {
    "max_fetch_mb": 100,
    "llm": {
        "context_length": 2048,
        "temperature": 0.4,
        "top_p": 0.9,
        "repeat_penalty": 1.1,
        "max_tokens": 96,
    },
    "job_concurrency": 1,
    "image_engine": {"type": "None", "url": "http://127.0.0.1:8188", "workflow_path": "/data/comfyui_workflow.json"},
}

DEFAULT_PRESETS = [
    {
        "name": "Default Meme",
        "selected_tones": ["funny"],
        "default_word_count": 12,
        "default_line_count": 2,
        "default_size_preset": "1080x1080",
        "default_background_mode": "solid",
        "default_text_template": "top_bottom",
        "default_model": "",
        "default_llm_parameters": DEFAULT_APP_SETTINGS["llm"],
    }
]

SUPPORTED_TEXT = {".txt", ".md", ".html", ".json", ".csv", ".pdf", ".docx"}
SUPPORTED_MODELS = {".gguf"}


def load_json(path: Path, default: Any) -> Any:
    if not path.exists():
        path.write_text(json.dumps(default, indent=2), encoding="utf-8")
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def save_json(path: Path, payload: Any) -> None:
    path.write_bytes(orjson.dumps(payload, option=orjson.OPT_INDENT_2))


APP_SETTINGS_FILE = DATA_DIR / "settings.json"
TONES_FILE = DATA_DIR / "tones.json"
PRESETS_FILE = DATA_DIR / "presets.json"
JOBS_FILE = DATA_DIR / "jobs.json"
WORKFLOW_PLACEHOLDER = DATA_DIR / "comfyui_workflow.json"

load_json(APP_SETTINGS_FILE, DEFAULT_APP_SETTINGS)
load_json(TONES_FILE, DEFAULT_TONES)
load_json(PRESETS_FILE, DEFAULT_PRESETS)
load_json(JOBS_FILE, {})
if not WORKFLOW_PLACEHOLDER.exists():
    WORKFLOW_PLACEHOLDER.write_text('{"note":"Replace with valid ComfyUI workflow JSON"}', encoding="utf-8")


class ExtractRequest(BaseModel):
    source_mode: Literal["url", "file", "both"]
    url: str | None = None
    max_fetch_mb_override: int | None = None


class TextGenRequest(BaseModel):
    content: str
    model: str
    tones: list[str]
    word_count_target: int = 12
    line_count_target: int = 2
    variants_per_tone: int = 1
    llm_parameters: dict[str, Any] = Field(default_factory=dict)


class RenderRequest(BaseModel):
    text_blocks: list[str]
    tone: str
    source_type: str
    source_reference: str
    model_used: str
    prompt_used: str
    size_preset: str = "1080x1080"
    custom_width: int | None = None
    custom_height: int | None = None
    export_jpg: bool = False
    background_mode: Literal["solid", "upload", "ai"] = "solid"
    solid_color: str = "#202020"
    template: Literal["top_bottom", "center_caption", "three_panel", "freeform_two"] = "top_bottom"
    font_name: str = "DejaVuSans.ttf"
    font_size: int = 56
    bold: bool = False
    letter_spacing: int = 0
    line_spacing: int = 8
    alignment: Literal["left", "center", "right"] = "center"
    outline: bool = True
    outline_thickness: int = 3
    shadow: bool = True


job_queue: asyncio.Queue = asyncio.Queue()
job_status: dict[str, dict[str, Any]] = load_json(JOBS_FILE, {})


app = FastAPI(title="Meme Generator", version="1.0.0")
app.mount("/static", StaticFiles(directory="app/static"), name="static")
app.mount("/outputs", StaticFiles(directory=str(OUTPUT_DIR)), name="outputs")
templates = Jinja2Templates(directory="app/templates")


def scan_models() -> list[dict[str, str]]:
    models: list[dict[str, str]] = []
    for path in MODELS_DIR.glob("**/*"):
        if path.is_file():
            ext = path.suffix.lower()
            models.append({
                "name": path.name,
                "path": str(path),
                "supported": "yes" if ext in SUPPORTED_MODELS else "no",
            })
    return sorted(models, key=lambda x: x["name"].lower())


def available_fonts() -> list[str]:
    bundled = ["DejaVuSans.ttf", "DejaVuSans-Bold.ttf"]
    custom = [p.name for p in FONTS_DIR.glob("*.ttf")]
    return sorted(set(bundled + custom))


def load_font(name: str, size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        FONTS_DIR / name,
        Path("/usr/share/fonts/truetype/dejavu") / ("DejaVuSans-Bold.ttf" if bold else name),
        Path("/usr/share/fonts/truetype/dejavu") / "DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


async def fetch_url_text(url: str, max_mb: int) -> str:
    max_bytes = max_mb * 1024 * 1024
    timeout = httpx.Timeout(15.0, connect=8.0)
    async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as client:
        r = await client.get(url)
        r.raise_for_status()
        raw = r.content[:max_bytes]
    doc = ReadabilityDocument(raw)
    html = doc.summary()
    soup = BeautifulSoup(html, "html.parser")
    text = "\n".join(s.strip() for s in soup.stripped_strings)
    return text


def parse_file(file_name: str, data: bytes) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_TEXT:
        return f"[Unsupported file format: {file_name}]"
    if suffix in {".txt", ".md", ".html", ".json"}:
        return data.decode("utf-8", errors="ignore")
    if suffix == ".csv":
        reader = csv.reader(data.decode("utf-8", errors="ignore").splitlines())
        return "\n".join([", ".join(r) for r in reader])
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        doc = Document(BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    return ""


def clean_excerpt(content: str, char_limit: int = 10000) -> str:
    return content.strip()[:char_limit]


def build_prompt(tone_template: str, content: str, word_count: int, line_count: int) -> str:
    return (
        f"{tone_template}\n"
        f"Target words: {word_count}. Target lines: {line_count}.\n"
        "Return only the final meme text with line breaks and no explanations.\n"
        f"Source content:\n{content[:7000]}"
    )


def call_llm(model_path: str, prompt: str, llm_params: dict[str, Any]) -> str:
    try:
        from llama_cpp import Llama
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"llama-cpp-python unavailable: {exc}") from exc
    if not Path(model_path).exists():
        raise HTTPException(status_code=400, detail="Selected model path does not exist")
    if Path(model_path).suffix.lower() != ".gguf":
        raise HTTPException(status_code=400, detail="Selected model is unsupported; only .gguf can be loaded")

    n_gpu_layers = -1 if os.getenv("USE_GPU", "1") == "1" else 0
    llm = Llama(
        model_path=model_path,
        n_ctx=int(llm_params.get("context_length", 2048)),
        n_gpu_layers=n_gpu_layers,
        verbose=False,
    )
    out = llm(
        prompt,
        max_tokens=int(llm_params.get("max_tokens", 96)),
        temperature=float(llm_params.get("temperature", 0.4)),
        top_p=float(llm_params.get("top_p", 0.9)),
        repeat_penalty=float(llm_params.get("repeat_penalty", 1.1)),
        stop=["\n\n"],
    )
    return out["choices"][0]["text"].strip()


def parse_size(req: RenderRequest) -> tuple[int, int]:
    presets = {
        "1080x1080": (1080, 1080),
        "1920x1080": (1920, 1080),
        "1080x1920": (1080, 1920),
        "1200x628": (1200, 628),
    }
    if req.size_preset == "custom":
        return req.custom_width or 1080, req.custom_height or 1080
    return presets.get(req.size_preset, (1080, 1080))


def fit_text(draw: ImageDraw.ImageDraw, text: str, box: tuple[int, int, int, int], font_name: str, base_size: int, bold: bool, line_spacing: int) -> tuple[ImageFont.FreeTypeFont, list[str], bool]:
    x0, y0, x1, y1 = box
    max_w = x1 - x0
    max_h = y1 - y0
    warning = False
    for size in range(base_size, 11, -2):
        font = load_font(font_name, size, bold=bold)
        lines: list[str] = []
        for paragraph in text.split("\n"):
            wrapped = textwrap.wrap(paragraph, width=max(10, int(max_w / max(size, 1) * 2.0))) or [""]
            lines.extend(wrapped)
        total_h = sum(draw.textbbox((0, 0), ln, font=font)[3] for ln in lines) + line_spacing * max(0, len(lines) - 1)
        max_line_w = max((draw.textbbox((0, 0), ln, font=font)[2] for ln in lines), default=0)
        if total_h <= max_h and max_line_w <= max_w:
            return font, lines, warning
    warning = True
    font = load_font(font_name, 12, bold=bold)
    lines = textwrap.wrap(text, width=20)
    return font, lines, warning


def draw_block(img: Image.Image, text: str, box: tuple[int, int, int, int], req: RenderRequest) -> bool:
    draw = ImageDraw.Draw(img)
    font, lines, warning = fit_text(draw, text, box, req.font_name, req.font_size, req.bold, req.line_spacing)
    x0, y0, x1, y1 = box
    block_h = sum(draw.textbbox((0, 0), ln, font=font)[3] for ln in lines) + req.line_spacing * max(0, len(lines)-1)
    y = y0 + max(0, (y1 - y0 - block_h)//2)
    for line in lines:
        width = draw.textbbox((0, 0), line, font=font)[2]
        if req.alignment == "left":
            x = x0
        elif req.alignment == "right":
            x = x1 - width
        else:
            x = x0 + max(0, (x1 - x0 - width)//2)
        if req.shadow:
            draw.text((x+2, y+2), line, font=font, fill=(0, 0, 0, 180))
        if req.outline:
            draw.text((x, y), line, font=font, fill="white", stroke_width=req.outline_thickness, stroke_fill="black")
        else:
            draw.text((x, y), line, font=font, fill="white")
        y += draw.textbbox((0, 0), line, font=font)[3] + req.line_spacing
    return warning


async def generate_ai_background(width: int, height: int) -> Image.Image:
    app_cfg = load_json(APP_SETTINGS_FILE, DEFAULT_APP_SETTINGS)
    engine = app_cfg.get("image_engine", {})
    if engine.get("type") != "ComfyUI":
        raise HTTPException(status_code=400, detail="AI background engine is not enabled")
    payload = load_json(Path(engine.get("workflow_path", "/data/comfyui_workflow.json")), {})
    async with httpx.AsyncClient(timeout=20.0) as client:
        r = await client.post(f"{engine.get('url', '').rstrip('/')}/prompt", json=payload)
    if r.status_code >= 400:
        raise HTTPException(status_code=502, detail="ComfyUI endpoint unreachable")
    return Image.new("RGB", (width, height), "#444")


def render_meme_image(req: RenderRequest, uploaded_bg: bytes | None = None) -> tuple[list[Path], bool]:
    width, height = parse_size(req)
    if req.background_mode == "solid":
        img = Image.new("RGB", (width, height), ImageColor.getrgb(req.solid_color))
    elif req.background_mode == "upload" and uploaded_bg:
        bg = Image.open(BytesIO(uploaded_bg)).convert("RGB")
        img = ImageEnhance.Contrast(bg.resize((width, height))).enhance(1.05)
    else:
        img = Image.new("RGB", (width, height), "#303030")

    boxes = {
        "top_bottom": [(30, 30, width - 30, height // 2 - 10), (30, height // 2 + 10, width - 30, height - 30)],
        "center_caption": [(60, height // 3, width - 60, 2 * height // 3)],
        "three_panel": [(40, 30, width - 40, height // 3), (40, height // 3, width - 40, 2 * height // 3), (40, 2 * height // 3, width - 40, height - 30)],
        "freeform_two": [(40, 40, width - 40, height // 2), (40, height // 2, width - 40, height - 40)],
    }
    warning = False
    for i, block in enumerate(req.text_blocks):
        if i >= len(boxes[req.template]):
            break
        warning = draw_block(img, block, boxes[req.template][i], req) or warning

    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    base = f"meme_{stamp}_{uuid4().hex[:8]}"
    png_path = OUTPUT_DIR / f"{base}.png"
    img.save(png_path, "PNG", optimize=True)
    paths = [png_path]
    if req.export_jpg:
        jpg_path = OUTPUT_DIR / f"{base}.jpg"
        img.save(jpg_path, "JPEG", quality=92)
        paths.append(jpg_path)

    meta = {
        "created_at": datetime.utcnow().isoformat(),
        "source_type": req.source_type,
        "source_reference": req.source_reference,
        "model_used": req.model_used,
        "tone": req.tone,
        "prompt_used": req.prompt_used,
        "size": f"{width}x{height}",
        "rendering_options": req.model_dump(),
        "overflow_warning": warning,
    }
    (OUTPUT_DIR / f"{base}.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    return paths, warning


@app.on_event("startup")
async def startup_event() -> None:
    async def worker() -> None:
        while True:
            job_id, req = await job_queue.get()
            try:
                job_status[job_id]["status"] = "running"
                save_json(JOBS_FILE, job_status)
                result = call_llm(req["model"], req["prompt"], req["llm_parameters"])
                job_status[job_id]["status"] = "done"
                job_status[job_id]["result"] = result
            except Exception as exc:
                job_status[job_id]["status"] = "error"
                job_status[job_id]["error"] = str(exc)
            finally:
                save_json(JOBS_FILE, job_status)
                job_queue.task_done()

    app_cfg = load_json(APP_SETTINGS_FILE, DEFAULT_APP_SETTINGS)
    concurrency = max(1, int(app_cfg.get("job_concurrency", 1)))
    for _ in range(concurrency):
        asyncio.create_task(worker())


@app.get("/", response_class=HTMLResponse)
async def home(request: Request) -> HTMLResponse:
    return templates.TemplateResponse("index.html", {"request": request})


@app.post("/api/extract")
async def extract_content(
    source_mode: str = Form(...),
    url: str = Form(""),
    excerpt_chars: int = Form(10000),
    max_fetch_mb_override: int = Form(0),
    files: list[UploadFile] = File(default=[]),
) -> dict[str, Any]:
    app_cfg = load_json(APP_SETTINGS_FILE, DEFAULT_APP_SETTINGS)
    max_mb = max_fetch_mb_override or app_cfg.get("max_fetch_mb", settings.max_fetch_mb)
    texts: list[str] = []
    refs: list[str] = []
    if source_mode in {"url", "both"} and url.strip():
        texts.append(await fetch_url_text(url.strip(), max_mb=max_mb))
        refs.append(url.strip())
    if source_mode in {"file", "both"}:
        for file in files:
            data = await file.read()
            texts.append(parse_file(file.filename, data))
            refs.append(file.filename)
    content = clean_excerpt("\n\n".join(texts), excerpt_chars)
    return {"text": content, "chars": len(content), "words": len(content.split()), "source_reference": ",".join(refs)}


@app.get("/api/models")
async def models() -> dict[str, Any]:
    return {"models": scan_models()}


@app.post("/api/generate_text")
async def generate_text(req: TextGenRequest) -> dict[str, Any]:
    tones = load_json(TONES_FILE, DEFAULT_TONES)
    if not req.tones:
        raise HTTPException(status_code=400, detail="At least one tone is required")
    variants = []
    for tone_name in req.tones:
        tone_template = tones.get(tone_name)
        if not tone_template:
            continue
        for _ in range(max(1, req.variants_per_tone)):
            prompt = build_prompt(tone_template, req.content, req.word_count_target, req.line_count_target)
            job_id = uuid4().hex
            job_status[job_id] = {"status": "queued", "tone": tone_name}
            await job_queue.put((job_id, {"model": req.model, "prompt": prompt, "llm_parameters": req.llm_parameters or load_json(APP_SETTINGS_FILE, DEFAULT_APP_SETTINGS)["llm"]}))
            variants.append({"job_id": job_id, "tone": tone_name, "prompt": prompt})
    save_json(JOBS_FILE, job_status)
    return {"jobs": variants}


@app.get("/api/jobs/{job_id}")
async def get_job(job_id: str) -> dict[str, Any]:
    item = job_status.get(job_id)
    if not item:
        raise HTTPException(status_code=404, detail="Job not found")
    return item


@app.post("/api/render_meme")
async def render_meme(request: Request, payload: str = Form(...), background_file: UploadFile | None = File(default=None)) -> dict[str, Any]:
    req = RenderRequest(**json.loads(payload))
    bg = await background_file.read() if background_file else None
    if req.background_mode == "ai":
        width, height = parse_size(req)
        img = await generate_ai_background(width, height)
        bio = BytesIO()
        img.save(bio, format="PNG")
        bg = bio.getvalue()
        req.background_mode = "upload"
    paths, warning = render_meme_image(req, uploaded_bg=bg)
    return {"files": [str(p) for p in paths], "overflow_warning": warning}


@app.get("/api/gallery")
async def gallery(tone: str = "", size: str = "") -> dict[str, Any]:
    items = []
    for image_path in sorted(OUTPUT_DIR.glob("*.png"), reverse=True):
        meta_path = image_path.with_suffix(".json")
        meta = json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.exists() else {}
        if tone and meta.get("tone") != tone:
            continue
        if size and meta.get("size") != size:
            continue
        items.append({"image": str(image_path), "meta": meta})
    return {"items": items}


@app.get("/api/settings")
async def get_settings() -> dict[str, Any]:
    return {
        "settings": load_json(APP_SETTINGS_FILE, DEFAULT_APP_SETTINGS),
        "tones": load_json(TONES_FILE, DEFAULT_TONES),
        "fonts": available_fonts(),
    }


@app.post("/api/settings")
async def update_settings(payload: dict[str, Any]) -> dict[str, str]:
    save_json(APP_SETTINGS_FILE, payload)
    return {"status": "ok"}


@app.post("/api/tones")
async def update_tones(payload: dict[str, str]) -> dict[str, str]:
    save_json(TONES_FILE, payload)
    return {"status": "ok"}


@app.get("/api/presets")
async def get_presets() -> dict[str, Any]:
    return {"presets": load_json(PRESETS_FILE, DEFAULT_PRESETS)}


@app.post("/api/presets")
async def set_presets(payload: list[dict[str, Any]]) -> dict[str, str]:
    save_json(PRESETS_FILE, payload)
    return {"status": "ok"}
